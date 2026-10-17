"""XLSX adapter (D-16 §3, D-17).

* Only `.xlsx`, detected from content: legacy `.xls` (OLE2) and macro-enabled workbooks
  (`.xlsm` content type or a VBA project) are refused, whatever the file is called.
* Zip-bomb guard before anything is decompressed: total uncompressed size and the total
  compression ratio are bounded.
* openpyxl in `read_only=True, data_only=True` mode with defusedxml installed. Formulas are
  never evaluated: a formula cell yields its cached value and a flag (a second, formula-mode
  pass over the same rows tells which cells hold formulas).
* Rows come from openpyxl's sheet parser directly (`_sheet_rows`), not from `iter_rows`,
  which yields one padded row per missing row number. A row number beyond Excel's
  1,048,576 rows is an invalid workbook, and a gap of missing rows is one blank draft
  whatever its length (Phase 8 review D2).
* openpyxl's warnings never leave the process: they can quote a cell value (an
  out-of-range date serial) and would bypass the log redaction (Phase 8 review D3).
* XLSX stores numbers as IEEE doubles: a numeric cell becomes `Decimal(repr(x))` (the
  shortest round trip), flagged beyond 15 significant digits. Dates become ISO strings.
* One sheet: the profile's named sheet, else the first. The marker is the sheet name
  `ILLUSTRATIVE SYNTHETIC` or a workbook title starting with the full marker (D-24 §2).
* Empty cells inside the header width read as ""; cells beyond it are kept as extras.
"""

import io
import warnings
import zipfile
from collections.abc import Iterator
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.cell.read_only import EMPTY_CELL, ReadOnlyCell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet._reader import WorkSheetParser

from app.ingestion.adapters.base import (
    ZIP_MAGIC,
    AdapterError,
    AdapterOptions,
    Draft,
    Limits,
    Locator,
    RowFlag,
    RowStream,
    SourceAdapter,
    check_header,
    from_cells,
    too_many_columns,
)
from app.ingestion.constants import (
    SYNTHETIC_MARKER,
    XLSX_MARKER_SHEET,
    RowFlagCode,
    SourceFormat,
    UploadRefusal,
)

_PARTS = ("[Content_Types].xml", "xl/workbook.xml")
_MACRO_TYPE = b"macroEnabled"
_EXACT_DIGITS = 15  # a double holds 15 significant decimal digits exactly
MAX_SHEET_ROWS = 1_048_576  # Excel's last row; a row number beyond it is not a workbook's
OPENPYXL_MODULES = r"openpyxl(\..*)?\Z"


def silence_openpyxl_warnings() -> None:
    """Phase 8 review D3: drop every warning raised from openpyxl's modules. They go to raw
    stderr, past the log redaction, and one quotes the serial value of an out-of-range date.
    The filter is process-wide and installed when this module is imported: parsing runs in
    a thread pool, where `warnings.catch_warnings` is not safe. Under pytest, the
    `filterwarnings` setting in pyproject.toml makes the same warnings errors instead."""
    warnings.filterwarnings("ignore", module=OPENPYXL_MODULES)


silence_openpyxl_warnings()


def _invalid() -> AdapterError:
    return AdapterError(UploadRefusal.XLSX_INVALID, "The workbook could not be read.")


def _significant_digits(value: Decimal) -> int:
    return len(value.normalize().as_tuple().digits)


def _text(cell: Any) -> tuple[str | None, list[RowFlagCode]]:
    """A cell's value as text, plus any flag the conversion raises."""
    value = cell.value
    flags: list[RowFlagCode] = []
    if getattr(cell, "data_type", None) == "e":
        flags.append(RowFlagCode.CELL_ERROR)
    if value is None:
        return None, flags
    if isinstance(value, bool):
        return ("true" if value else "false"), flags
    if isinstance(value, int | float):
        number = Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
        if number.is_finite() and _significant_digits(number) > _EXACT_DIGITS:
            flags.append(RowFlagCode.PRECISION_BEYOND_15_DIGITS)
        return str(number), flags
    if isinstance(value, datetime):
        if value.time() == time(0) and value.tzinfo is None:
            return value.date().isoformat(), flags
        return value.isoformat(), flags
    if isinstance(value, date | time):
        return value.isoformat(), flags
    return str(value), flags


Row = tuple[int, tuple[Any, ...]]  # (row number, cells)


def _sheet_rows(ws: ReadOnlyWorksheet, limits: Limits) -> Iterator[Row]:
    """(row number, cells) for each <row> element of the sheet, read through openpyxl's own
    sheet parser (private API, pinned openpyxl <3.2), with the same cell conversion as
    `iter_rows`. `iter_rows` is not used because it pads: one empty row per missing row
    number (`<row r="2000000000">` is two billion rows) and every row up to its last
    styled cell. Here:

    * a row number beyond MAX_SHEET_ROWS is `xlsx_invalid`, before any padding;
    * missing row numbers are left to the caller, as a gap;
    * a row's width counts only the cells holding a value, refused above the column limit.

    As in `iter_rows`, a row whose number does not increase is skipped (Phase 8 review D4,
    an open follow-up)."""
    sheet: Any = ws  # the stubs model Worksheet, not the read-only internals
    wb: Any = ws.parent
    with sheet._get_source() as src:
        parser = WorkSheetParser(
            src,
            sheet._shared_strings,
            data_only=wb.data_only,
            epoch=wb.epoch,
            date_formats=wb._date_formats,
            timedelta_formats=wb._timedelta_formats,
        )
        previous = 0
        for number, row in parser.parse():
            if number > MAX_SHEET_ROWS:
                raise _invalid()
            if number <= previous:
                continue
            previous = number
            width = max((c["column"] for c in row if c["value"] is not None), default=0)
            if width > limits.max_columns:
                raise too_many_columns(limits, {"sheet": ws.title, "row": number})
            cells: list[Any] = [EMPTY_CELL] * width
            for c in row:
                if c["column"] <= width:
                    cells[c["column"] - 1] = ReadOnlyCell(ws, **c)
            yield number, tuple(cells)


class XlsxAdapter(SourceAdapter):
    format = SourceFormat.XLSX
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def detect(self, data: bytes) -> bool:
        if not data.startswith(ZIP_MAGIC):
            return False
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
        except (zipfile.BadZipFile, ValueError, OSError):
            return False
        return all(p in names for p in _PARTS)

    def iter_rows(self, data: bytes, options: AdapterOptions, limits: Limits) -> RowStream:
        _check_archive(data, limits)
        try:
            values = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
            formulas = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
        except Exception:  # any parser failure: XML, zip, missing part, entity attack
            raise _invalid() from None
        try:
            ws_v, ws_f = _sheet(values, options), _sheet(formulas, options)
            rows = _paired(_sheet_rows(ws_v, limits), _sheet_rows(ws_f, limits))
            header = next(rows, None)
            title = values.properties.title or ""
            marker = ws_v.title.startswith(XLSX_MARKER_SHEET) or title.startswith(SYNTHETIC_MARKER)
            columns = _header(header, ws_v.title, limits)
        except BaseException:
            values.close()
            formulas.close()
            raise
        sheet = ws_v.title
        return RowStream(
            columns, marker, _drafts(rows, columns, sheet, limits, values, formulas), sheet
        )


def _check_archive(data: bytes, limits: Limits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            total = sum(i.file_size for i in infos)
            packed = sum(i.compress_size for i in infos)
            if total > limits.max_xlsx_uncompressed_bytes:
                raise AdapterError(
                    UploadRefusal.XLSX_UNCOMPRESSED_TOO_LARGE,
                    "The workbook is too large once uncompressed.",
                    {"limit": limits.max_xlsx_uncompressed_bytes},
                )
            if total > limits.max_xlsx_ratio * max(packed, 1):
                raise AdapterError(
                    UploadRefusal.XLSX_COMPRESSION_RATIO_TOO_HIGH,
                    "The workbook's compression ratio is too high.",
                    {"limit": limits.max_xlsx_ratio},
                )
            names = {i.filename for i in infos}
            if "xl/vbaProject.bin" in names or _MACRO_TYPE in zf.read("[Content_Types].xml"):
                raise AdapterError(
                    UploadRefusal.MACRO_ENABLED_WORKBOOK,
                    "Macro-enabled workbooks (.xlsm) are not accepted; save it as .xlsx.",
                )
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise _invalid() from None


def _sheet(wb: Workbook, options: AdapterOptions) -> ReadOnlyWorksheet:
    ws: Any
    if options.sheet_name is not None:
        if options.sheet_name not in wb.sheetnames:
            raise AdapterError(
                UploadRefusal.SHEET_NOT_FOUND, "The workbook has no sheet with the profile's name."
            )
        ws = wb[options.sheet_name]
    else:
        if not wb.worksheets:
            raise AdapterError(UploadRefusal.MISSING_HEADER, "The workbook has no sheet.")
        ws = wb.worksheets[0]
    if not isinstance(ws, ReadOnlyWorksheet):
        raise _invalid()
    return ws


def _paired(values: Iterator[Row], formulas: Iterator[Row]) -> Iterator[tuple[Row, Row]]:
    """Each row from both passes. A parser failure mid-sheet is an invalid workbook."""
    try:
        for row_v, row_f in zip(values, formulas, strict=True):
            if row_v[0] != row_f[0]:
                raise _invalid()
            yield row_v, row_f
    except AdapterError:
        raise
    except Exception:
        raise _invalid() from None


def _header(header: tuple[Row, Row] | None, sheet: str, limits: Limits) -> list[str]:
    """Row 1. A sheet whose first row is a later one has a blank header row."""
    cells = [_text(c)[0] for c in header[0][1]] if header and header[0][0] == 1 else []
    while cells and cells[-1] is None:
        cells.pop()
    if not cells:
        raise AdapterError(UploadRefusal.MISSING_HEADER, "The sheet has no header row.")
    return check_header([c or "" for c in cells], {"sheet": sheet, "row": 1}, limits)


def _drafts(
    rows: Iterator[tuple[Row, Row]],
    columns: list[str],
    sheet: str,
    limits: Limits,
    *workbooks: Workbook,
) -> Iterator[Draft]:
    try:
        expected = 2  # the row after the header
        for (number, row_v), (_, row_f) in rows:
            if number > expected:  # missing row numbers: blank rows, as one draft
                gap: Locator = {"sheet": sheet, "row": expected}
                yield Draft(gap, {}, blank=True, count=number - expected)
            expected = number + 1
            locator: Locator = {"sheet": sheet, "row": number}
            texts: list[str | None] = []
            flags: list[RowFlag] = []
            # A formula without a cached value has no value in the values pass, so walk
            # the wider of the two passes.
            for i in range(1, max(len(row_v), len(row_f)) + 1):
                text, codes = _text(row_v[i - 1] if i <= len(row_v) else EMPTY_CELL)
                formula = i <= len(row_f) and getattr(row_f[i - 1], "data_type", None) == "f"
                if formula:
                    codes.append(
                        RowFlagCode.FORMULA_CACHED_VALUE
                        if text is not None
                        else RowFlagCode.FORMULA_WITHOUT_CACHED_VALUE
                    )
                flags.extend(RowFlag(code, i) for code in codes)
                texts.append(text)
            while texts and texts[-1] is None:
                texts.pop()
            if not texts and not flags:
                yield Draft(locator, {}, blank=True)
                continue
            cells = [t or "" for t in texts]
            if len(cells) < len(columns):
                cells += [""] * (len(columns) - len(cells))  # an empty cell is ""
            yield from_cells(columns, cells, locator, limits, flags, flag_missing=False)
    finally:
        for wb in workbooks:
            wb.close()
